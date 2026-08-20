from __future__ import annotations

import unittest
from io import BytesIO

from openpyxl import Workbook

from app.eitan_library import eitan_search_groups, parse_eitan_library_upload
from app.html_layout_engine import _public_categories
from app.bulletin_models import (
    BulletinCategory,
    BulletinData,
    BulletinMeta,
    BulletinPerson,
    BulletinStatement,
)
from app.weekday_art import weekday_cover_png_bytes, weekday_cover_svg
from app.weekday_palette import weekday_palette_for_report
from app.word_rtl import apply_persian_document, rtl_table


def _xlsx(rows: list[list[object]], sheet: str = "کتابخانه") -> bytes:
    workbook = Workbook()
    ws = workbook.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class EitanLibraryTests(unittest.TestCase):
    def test_paired_keyword_and_person_become_and_groups(self) -> None:
        payload = _xlsx(
            [
                ["کلیدواژه", "فرد", "دسته", "وزن"],
                ["تنگه هرمز", "عباس عراقچی", "امنیت دریایی", 3],
                ["نفتکش", "حسین سلامی", "امنیت دریایی", 2],
                ["خاموشی", "", "برق", 1],
            ]
        )
        library = parse_eitan_library_upload("library.xlsx", payload)
        self.assertTrue(library.paired)
        self.assertIn("تنگه هرمز", library.keywords)
        self.assertIn("عباس عراقچی", library.people)
        self.assertEqual(library.clusters[0]["name"], "امنیت دریایی")
        groups = eitan_search_groups(library)
        self.assertIn(["تنگه هرمز", "عباس عراقچی"], groups)
        self.assertIn(["خاموشی"], groups)

    def test_single_column_stays_or_search(self) -> None:
        payload = _xlsx([["کلیدواژه"], ["هرمز"], ["نفتکش"]])
        library = parse_eitan_library_upload("keywords.xlsx", payload)
        self.assertFalse(library.paired)
        self.assertEqual(eitan_search_groups(library), [["هرمز"], ["نفتکش"]])


class LayoutOrderTests(unittest.TestCase):
    def test_events_come_last(self) -> None:
        meta = BulletinMeta(
            issue_number=1,
            report_mode="concise",
            report_date_jalali="1405/05/26",
            generated_at="2026-08-17T00:00:00+00:00",
        )
        statement = BulletinStatement(
            statement_id="s1",
            topic="آزمایش",
            statement_mode="media_report",
            summary_short="کوتاه",
            summary_detailed="بلند",
            importance_score=0.5,
            include_in_main=True,
        )
        person = BulletinPerson(name_canonical="شخص", is_registry_person=True, statements=[statement])
        data = BulletinData(
            meta=meta,
            categories=[
                BulletinCategory(category_id="events", title="وقایع و رویدادهای مهم ایران و جهان", order=1, section_letter="الف", people=[person]),
                BulletinCategory(category_id="government", title="مسئولان دولت", order=2, section_letter="ب", people=[person]),
            ],
        )
        keys = [category.category_id for category, _people in _public_categories(data)]
        self.assertEqual(keys[-1], "events")
        self.assertEqual(keys[0], "government")


class WeekdayArtTests(unittest.TestCase):
    def test_svg_and_png_for_palette(self) -> None:
        palette = weekday_palette_for_report("1405/05/26")
        svg = weekday_cover_svg(palette)
        self.assertIn(str(palette.get("name") or ""), svg)
        png = weekday_cover_png_bytes(palette)
        self.assertGreater(len(png.getvalue()), 100)


class WordRtlTests(unittest.TestCase):
    def test_document_language_is_persian(self) -> None:
        from docx import Document
        from docx.oxml.ns import qn

        doc = Document()
        apply_persian_document(doc)
        table = doc.add_table(rows=1, cols=1)
        rtl_table(table)
        settings = doc.settings.element.find(qn("w:themeFontLang"))
        self.assertIsNotNone(settings)
        self.assertEqual(settings.get(qn("w:bidi")), "fa-IR")
        self.assertIsNotNone(table._tbl.tblPr.find(qn("w:bidiVisual")))


if __name__ == "__main__":
    unittest.main()
