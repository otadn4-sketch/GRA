from __future__ import annotations

"""Parse Eitan Gara Excel libraries and turn rows into search groups and charts."""

import csv
import io
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from typing import Any, Iterable

from .persian_text import canonical_key, normalize_persian


EITAN_FILE_MAX_BYTES = 5 * 1024 * 1024
EITAN_TEXT_MAX_CHARS = 800_000
EITAN_SEARCH_GROUP_LIMIT = 120
EITAN_SEARCH_TERM_LIMIT = 480
EITAN_AND_WIDTH = 4
EITAN_INSIGHT_SAMPLE_LIMIT = 1200

_QUERY_STOPWORDS = {"or", "and", "not", "near", "یا", "و", "نه"}

_EITAN_HEADER_TERMS = {
    "کلیدواژه",
    "کلیدواژه‌ها",
    "کلید واژه",
    "keyword",
    "keywords",
    "نام",
    "افراد",
    "شخص",
    "افراد شاخص",
    "title",
    "name",
    "دسته",
    "گروه",
    "خوشه",
    "محور",
    "وزن",
    "weight",
}

_KEYWORD_HEADERS = {
    "کلیدواژه",
    "کلیدواژه‌ها",
    "کلید واژه",
    "عبارت",
    "واژه",
    "keyword",
    "keywords",
    "term",
    "terms",
    "canonical_term",
    "canonical term",
}
_PERSON_HEADERS = {
    "فرد",
    "شخص",
    "اشخاص",
    "افراد",
    "افراد شاخص",
    "شخصیت",
    "نام",
    "people",
    "person",
    "persons",
    "name",
    "canonical_name",
    "canonical name",
}
_CLUSTER_HEADERS = {
    "دسته",
    "گروه",
    "خوشه",
    "محور",
    "موضوع",
    "cluster",
    "group",
    "category",
    "theme",
    "axis",
}
_WEIGHT_HEADERS = {"وزن", "weight", "امتیاز", "score", "priority"}
_MATCH_HEADERS = {"عملگر", "منطق", "match", "operator", "and", "or"}

_KEYWORD_SHEETS = {
    "keyword_library",
    "keyword library",
    "keywords",
    "کتابخانه کلیدواژه",
    "کلیدواژه",
    "کلیدواژه‌ها",
}
_PERSON_SHEETS = {
    "person_library",
    "person library",
    "people",
    "persons",
    "کتابخانه افراد",
    "افراد",
    "اشخاص",
}

_KEYWORD_ID_HEADERS = {"keyword_id", "keyword id", "id", "شناسه", "شناسه کلیدواژه"}
_ALIAS_HEADERS = {"aliases", "alias", "مترادف", "مترادف‌ها", "نام‌های دیگر"}
_SUBCATEGORY_HEADERS = {"subcategory", "sub category", "sub_category", "زیردسته", "زیر دسته", "زیرطبقه"}
_HELPER_HEADERS = {"helper_terms", "helper terms", "helpers", "عبارات کمکی", "کمکی"}
_SEARCH_QUERY_HEADERS = {"search_query", "search query", "query", "عبارت جستجو", "جستجو"}
_PERSON_ID_HEADERS = {"person_id", "person id", "id", "شناسه", "شناسه فرد"}
_ROLE_HEADERS = {"role", "نقش", "سمت", "عنوان شغلی"}
_INSTITUTION_HEADERS = {"institution", "سازمان", "نهاد", "موسسه", "مؤسسه"}
_VIEW_CLUSTER_HEADERS = {"view_cluster", "view cluster", "cluster", "خوشه", "خوشه دید"}
_SUBCLUSTER_HEADERS = {"subcluster", "sub_cluster", "sub cluster", "زیرخوشه", "زیر خوشه"}
_TOPIC_HEADERS = {"topic_keywords", "topic keywords", "topics", "موضوعات", "کلیدواژه موضوع"}
_SOURCE_SCOPE_HEADERS = {"source_scope", "source scope", "scope", "محدوده منبع", "دامنه منبع"}


def _clean_term(value: Any) -> str:
    text = " ".join(str(value or "").split())
    text = normalize_persian(text)
    return text.strip(" -_|،,;/")


def _header_key(value: Any) -> str:
    return canonical_key(_clean_term(value)) or _clean_term(value).casefold()


def _keys(*names: str) -> set[str]:
    return {_header_key(name) for name in names if name}


def split_eitan_terms(text: str | None) -> list[str]:
    seen: set[str] = set()
    terms: list[str] = []
    for raw_line in str(text or "").replace("\r", "\n").split("\n"):
        line = " ".join(raw_line.split()).strip()
        if not line or line.startswith("#"):
            continue
        chunks = re.split(r"[|،;]+", line)
        if len(chunks) == 1 and "," in line and "\t" not in line:
            chunks = [part.strip() for part in line.split(",")]
        for chunk in chunks:
            term = _clean_term(chunk)
            if len(term) < 2:
                continue
            key = canonical_key(term) or term
            if key in _EITAN_HEADER_TERMS or key in seen:
                continue
            seen.add(key)
            terms.append(term)
    return terms


def _unique_terms(values: Iterable[Any]) -> list[str]:
    seen: set[str] = set()
    terms: list[str] = []
    for value in values:
        term = _clean_term(value)
        if len(term) < 2:
            continue
        key = canonical_key(term) or term
        if key in _EITAN_HEADER_TERMS or key in seen:
            continue
        seen.add(key)
        terms.append(term)
    return terms


def _split_query(text: str | None) -> list[str]:
    raw = str(text or "").strip()
    if not raw:
        return []
    quoted: list[str] = []
    for match in re.finditer(r'"([^"]+)"|«([^»]+)»', raw):
        quoted.append(match.group(1) or match.group(2) or "")
    remainder = re.sub(r'"[^"]*"|«[^»]*»', " ", raw)
    remainder = remainder.replace("؛", " ").replace(";", " ").replace("|", " ")
    remainder = remainder.replace("،", " ").replace(",", " ").replace("/", " ")
    tokens: list[str] = []
    for piece in re.split(r"\s+", remainder):
        cleaned = _clean_term(piece)
        if not cleaned or cleaned.casefold() in _QUERY_STOPWORDS:
            continue
        tokens.append(cleaned)
    return _unique_terms([*quoted, *tokens])


def _priority(value: Any) -> int:
    text = str(value or "").strip().replace(",", ".")
    if not text:
        return 1
    try:
        number = int(float(text))
    except ValueError:
        return 1
    return max(1, min(number, 10))


def _cell(row: list[Any], index: int | None) -> str:
    if index is None or index < 0 or index >= len(row):
        return ""
    return _clean_term(row[index])


def _find_col(headers: dict[str, int], names: set[str]) -> int | None:
    wanted = _keys(*names)
    for key, index in headers.items():
        if key in wanted:
            return index
    return None


def _header_map(row: list[Any]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, cell in enumerate(row):
        key = _header_key(cell)
        if key and key not in mapping:
            mapping[key] = index
    return mapping


@dataclass
class KeywordRecord:
    keyword_id: str = ""
    canonical_term: str = ""
    aliases: list[str] = field(default_factory=list)
    category: str = ""
    subcategory: str = ""
    priority: int = 1
    helper_terms: list[str] = field(default_factory=list)
    search_query: str = ""

    def search_terms(self) -> list[str]:
        return _unique_terms(
            [
                self.canonical_term,
                *self.aliases,
                *self.helper_terms,
                *_split_query(self.search_query),
            ]
        )

    def to_dict(self) -> dict[str, Any]:
        return {
            "keyword_id": self.keyword_id,
            "canonical_term": self.canonical_term,
            "aliases": self.aliases,
            "category": self.category,
            "subcategory": self.subcategory,
            "priority": self.priority,
            "helper_terms": self.helper_terms,
            "search_query": self.search_query,
        }


@dataclass
class PersonRecord:
    person_id: str = ""
    canonical_name: str = ""
    aliases: list[str] = field(default_factory=list)
    role: str = ""
    institution: str = ""
    view_cluster: str = ""
    subcluster: str = ""
    topic_keywords: list[str] = field(default_factory=list)
    priority: int = 1
    source_scope: str = ""
    search_query: str = ""

    def search_terms(self) -> list[str]:
        return _unique_terms(
            [
                self.canonical_name,
                *self.aliases,
                *self.topic_keywords,
                *_split_query(self.search_query),
            ]
        )

    def to_dict(self) -> dict[str, Any]:
        return {
            "person_id": self.person_id,
            "canonical_name": self.canonical_name,
            "aliases": self.aliases,
            "role": self.role,
            "institution": self.institution,
            "view_cluster": self.view_cluster,
            "subcluster": self.subcluster,
            "topic_keywords": self.topic_keywords,
            "priority": self.priority,
            "source_scope": self.source_scope,
            "search_query": self.search_query,
        }


def keyword_from_dict(item: dict[str, Any] | None) -> KeywordRecord | None:
    payload = item or {}
    term = _clean_term(payload.get("canonical_term") or payload.get("term") or "")
    if not term:
        return None
    return KeywordRecord(
        keyword_id=_clean_term(payload.get("keyword_id") or ""),
        canonical_term=term,
        aliases=_unique_terms(payload.get("aliases") or []),
        category=_clean_term(payload.get("category") or ""),
        subcategory=_clean_term(payload.get("subcategory") or ""),
        priority=_priority(payload.get("priority") or 1),
        helper_terms=_unique_terms(payload.get("helper_terms") or []),
        search_query=_clean_term(payload.get("search_query") or ""),
    )


def person_from_dict(item: dict[str, Any] | None) -> PersonRecord | None:
    payload = item or {}
    name = _clean_term(payload.get("canonical_name") or payload.get("name") or "")
    if not name:
        return None
    return PersonRecord(
        person_id=_clean_term(payload.get("person_id") or ""),
        canonical_name=name,
        aliases=_unique_terms(payload.get("aliases") or []),
        role=_clean_term(payload.get("role") or ""),
        institution=_clean_term(payload.get("institution") or ""),
        view_cluster=_clean_term(payload.get("view_cluster") or ""),
        subcluster=_clean_term(payload.get("subcluster") or ""),
        topic_keywords=_unique_terms(payload.get("topic_keywords") or []),
        priority=_priority(payload.get("priority") or 1),
        source_scope=_clean_term(payload.get("source_scope") or ""),
        search_query=_clean_term(payload.get("search_query") or ""),
    )


def _clusters_from_records(
    keyword_records: list[KeywordRecord],
    person_records: list[PersonRecord],
) -> list[dict[str, Any]]:
    buckets: dict[str, dict[str, list[str]]] = {}
    for record in keyword_records:
        name = record.category
        if not name:
            continue
        bucket = buckets.setdefault(name, {"keywords": [], "people": []})
        bucket["keywords"].extend(record.search_terms())
    for record in person_records:
        name = record.view_cluster or record.subcluster
        if not name:
            continue
        bucket = buckets.setdefault(name, {"keywords": [], "people": []})
        bucket["people"].extend(record.search_terms())
    return [
        {
            "name": name,
            "keywords": _unique_terms(values["keywords"]),
            "people": _unique_terms(values["people"]),
        }
        for name, values in buckets.items()
    ]


def _clauses_from_records(
    keyword_records: list[KeywordRecord],
    person_records: list[PersonRecord],
) -> list[dict[str, Any]]:
    clauses: list[dict[str, Any]] = []
    keywords_sorted = [item for _index, item in sorted(enumerate(keyword_records), key=lambda pair: (-pair[1].priority, pair[0]))]
    people_sorted = [item for _index, item in sorted(enumerate(person_records), key=lambda pair: (-pair[1].priority, pair[0]))]
    for record in keywords_sorted:
        terms = record.search_terms()
        if not terms:
            continue
        clauses.append(
            {
                "keywords": terms,
                "people": [],
                "cluster": record.category,
                "match": "any",
                "weight": float(record.priority),
            }
        )
    for record in people_sorted:
        terms = record.search_terms()
        if not terms:
            continue
        clauses.append(
            {
                "keywords": [],
                "people": terms,
                "cluster": record.view_cluster,
                "match": "any",
                "weight": float(record.priority),
            }
        )
    return clauses


@dataclass
class EitanLibrary:
    keywords: list[str] = field(default_factory=list)
    people: list[str] = field(default_factory=list)
    keyword_records: list[KeywordRecord] = field(default_factory=list)
    person_records: list[PersonRecord] = field(default_factory=list)
    clusters: list[dict[str, Any]] = field(default_factory=list)
    clauses: list[dict[str, Any]] = field(default_factory=list)
    filename: str = ""
    sheets: list[str] = field(default_factory=list)
    paired: bool = False
    version: int = 2

    def as_json(self) -> str:
        payload = {
            "version": self.version,
            "filename": self.filename,
            "sheets": self.sheets,
            "keywords": self.keywords,
            "people": self.people,
            "keyword_records": [item.to_dict() for item in self.keyword_records],
            "person_records": [item.to_dict() for item in self.person_records],
            "clusters": self.clusters,
            "clauses": self.clauses,
            "paired": self.paired,
        }
        text = json.dumps(payload, ensure_ascii=False)
        if len(text) > EITAN_TEXT_MAX_CHARS:
            raise ValueError("حجم کتابخانه بیش از حد مجاز است.")
        return text

    def summary(self) -> dict[str, Any]:
        categories = _unique_terms(item.category for item in self.keyword_records if item.category)
        view_clusters = _unique_terms(
            item.view_cluster for item in self.person_records if item.view_cluster
        )
        return {
            "filename": self.filename,
            "sheets": self.sheets,
            "keyword_count": len(self.keywords),
            "people_count": len(self.people),
            "cluster_count": len(self.clusters),
            "category_count": len(categories),
            "view_cluster_count": len(view_clusters),
            "clause_count": len(self.clauses),
            "paired": self.paired,
            "keyword_sample": self.keywords[:12],
            "people_sample": self.people[:12],
            "clusters": [item.get("name") for item in self.clusters[:12]],
            "categories": categories[:12],
        }


def _library_from_records(
    keyword_records: list[KeywordRecord],
    person_records: list[PersonRecord],
    *,
    filename: str = "",
    sheets: list[str] | None = None,
    paired: bool = False,
) -> EitanLibrary:
    keywords = _unique_terms(item.canonical_term for item in keyword_records)
    people = _unique_terms(item.canonical_name for item in person_records)
    return EitanLibrary(
        keywords=keywords,
        people=people,
        keyword_records=keyword_records,
        person_records=person_records,
        clusters=_clusters_from_records(keyword_records, person_records),
        clauses=_clauses_from_records(keyword_records, person_records),
        filename=filename,
        sheets=list(sheets or []),
        paired=paired,
        version=2,
    )


def library_from_parts(
    *,
    keywords: Iterable[str] | None = None,
    people: Iterable[str] | None = None,
    filename: str = "",
) -> EitanLibrary:
    keyword_records = [
        KeywordRecord(keyword_id=f"kw-{index}", canonical_term=term)
        for index, term in enumerate(_unique_terms(keywords or []), start=1)
    ]
    person_records = [
        PersonRecord(person_id=f"ps-{index}", canonical_name=name)
        for index, name in enumerate(_unique_terms(people or []), start=1)
    ]
    return _library_from_records(keyword_records, person_records, filename=filename, sheets=[])


def library_from_json(raw: str | None, *, keywords_text: str = "", people_text: str = "") -> EitanLibrary:
    text = str(raw or "").strip()
    if text:
        try:
            payload = json.loads(text)
        except json.JSONDecodeError:
            payload = None
        if isinstance(payload, dict):
            keyword_records = [
                record
                for record in (keyword_from_dict(item) for item in payload.get("keyword_records") or [])
                if record
            ]
            person_records = [
                record
                for record in (person_from_dict(item) for item in payload.get("person_records") or [])
                if record
            ]
            if not keyword_records and payload.get("keywords"):
                keyword_records = [
                    KeywordRecord(keyword_id=f"kw-{index}", canonical_term=term)
                    for index, term in enumerate(_unique_terms(payload.get("keywords") or []), start=1)
                ]
            if not person_records and payload.get("people"):
                person_records = [
                    PersonRecord(person_id=f"ps-{index}", canonical_name=name)
                    for index, name in enumerate(_unique_terms(payload.get("people") or []), start=1)
                ]
            if keyword_records or person_records:
                library = _library_from_records(
                    keyword_records,
                    person_records,
                    filename=str(payload.get("filename") or ""),
                    sheets=list(payload.get("sheets") or []),
                    paired=bool(payload.get("paired")),
                )
                if payload.get("clusters"):
                    library.clusters = list(payload.get("clusters") or [])
                if payload.get("clauses"):
                    library.clauses = list(payload.get("clauses") or [])
                return library
    return library_from_parts(
        keywords=split_eitan_terms(keywords_text),
        people=split_eitan_terms(people_text),
    )


def library_from_axis(axis: dict[str, Any] | None) -> EitanLibrary:
    row = axis or {}
    return library_from_json(
        row.get("library_json"),
        keywords_text=str(row.get("keywords_text") or ""),
        people_text=str(row.get("people_text") or ""),
    )


def _classify_header(header: str) -> str | None:
    key = _header_key(header)
    if key in _keys(*_KEYWORD_HEADERS):
        return "keyword"
    if key in _keys(*_PERSON_HEADERS):
        return "person"
    if key in _keys(*_CLUSTER_HEADERS):
        return "cluster"
    if key in _keys(*_WEIGHT_HEADERS):
        return "weight"
    if key in _keys(*_MATCH_HEADERS):
        return "match"
    return None


def _sheet_kind(name: str, headers: dict[str, int]) -> str:
    key = _header_key(name)
    compact = key.replace(" ", "_")
    if compact in _KEYWORD_SHEETS or key in _KEYWORD_SHEETS or "keyword" in compact:
        return "keywords"
    if compact in _PERSON_SHEETS or key in _PERSON_SHEETS or "person" in compact:
        return "people"
    header_keys = set(headers)
    if header_keys & _keys("canonical_term"):
        return "keywords"
    if header_keys & _keys("canonical_name"):
        return "people"
    if any(token in key for token in ("فرد", "شخص", "people", "person")):
        return "people"
    if any(token in key for token in ("کلید", "keyword", "واژه")):
        return "keywords"
    return "mixed"


def _parse_weight(value: Any) -> float:
    text = str(value or "").strip().replace(",", ".")
    if not text:
        return 1.0
    try:
        number = float(text)
    except ValueError:
        return 1.0
    if number <= 0:
        return 1.0
    return min(number, 100.0)


def _parse_match(value: Any, *, default: str) -> str:
    text = _header_key(value)
    if text in {"and", "و", "همه", "all"}:
        return "all"
    if text in {"or", "یا", "any", "هریک"}:
        return "any"
    return default


def _rows_from_xlsx(payload: bytes) -> tuple[list[str], list[tuple[str, list[list[Any]]]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    sheets: list[tuple[str, list[list[Any]]]] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[list[Any]] = []
            for row in sheet.iter_rows(min_row=1, max_col=24, max_row=8000, values_only=True):
                values = list(row)
                if not any(cell not in (None, "") for cell in values):
                    continue
                rows.append(values)
            sheets.append((str(sheet.title or "Sheet"), rows))
    finally:
        workbook.close()
    return [name for name, _rows in sheets], sheets


def _rows_from_csv(payload: bytes) -> tuple[list[str], list[tuple[str, list[list[Any]]]]]:
    sample = payload.decode("utf-8-sig", errors="replace")
    rows = [list(row) for row in csv.reader(io.StringIO(sample)) if any(str(cell or "").strip() for cell in row)]
    return ["csv"], [("csv", rows)]


def _rows_from_text(payload: bytes) -> tuple[list[str], list[tuple[str, list[list[Any]]]]]:
    lines = payload.decode("utf-8-sig", errors="replace").splitlines()
    rows = [[line] for line in lines if str(line or "").strip()]
    return ["txt"], [("txt", rows)]


def _parse_keyword_sheet(rows: list[list[Any]]) -> list[KeywordRecord]:
    if not rows:
        return []
    headers = _header_map(rows[0])
    term_col = _find_col(headers, _KEYWORD_HEADERS)
    if term_col is None:
        return []
    id_col = _find_col(headers, _KEYWORD_ID_HEADERS - {"id"}) or _find_col(headers, {"keyword_id", "keyword id"})
    alias_col = _find_col(headers, _ALIAS_HEADERS)
    cat_col = _find_col(headers, {"category", "دسته", "طبقه"})
    sub_col = _find_col(headers, _SUBCATEGORY_HEADERS)
    pri_col = _find_col(headers, _WEIGHT_HEADERS)
    helper_col = _find_col(headers, _HELPER_HEADERS)
    query_col = _find_col(headers, _SEARCH_QUERY_HEADERS)
    records: list[KeywordRecord] = []
    for index, row in enumerate(rows[1:], start=1):
        term = _cell(row, term_col)
        if not term:
            continue
        records.append(
            KeywordRecord(
                keyword_id=_cell(row, id_col) or f"kw-{index}",
                canonical_term=term,
                aliases=split_eitan_terms(_cell(row, alias_col)),
                category=_cell(row, cat_col),
                subcategory=_cell(row, sub_col),
                priority=_priority(_cell(row, pri_col) or 1),
                helper_terms=split_eitan_terms(_cell(row, helper_col)),
                search_query=_cell(row, query_col),
            )
        )
    return records


def _parse_person_sheet(rows: list[list[Any]]) -> list[PersonRecord]:
    if not rows:
        return []
    headers = _header_map(rows[0])
    name_col = _find_col(headers, _PERSON_HEADERS | {"canonical_name", "canonical name"})
    if name_col is None:
        name_col = _find_col(headers, {"canonical_name", "canonical name", "نام اصلی"})
    if name_col is None:
        return []
    id_col = _find_col(headers, {"person_id", "person id", "شناسه فرد"})
    alias_col = _find_col(headers, _ALIAS_HEADERS)
    role_col = _find_col(headers, _ROLE_HEADERS)
    inst_col = _find_col(headers, _INSTITUTION_HEADERS)
    cluster_col = _find_col(headers, _VIEW_CLUSTER_HEADERS)
    sub_col = _find_col(headers, _SUBCLUSTER_HEADERS)
    topic_col = _find_col(headers, _TOPIC_HEADERS)
    pri_col = _find_col(headers, _WEIGHT_HEADERS)
    scope_col = _find_col(headers, _SOURCE_SCOPE_HEADERS)
    query_col = _find_col(headers, _SEARCH_QUERY_HEADERS)
    records: list[PersonRecord] = []
    for index, row in enumerate(rows[1:], start=1):
        name = _cell(row, name_col)
        if not name:
            continue
        records.append(
            PersonRecord(
                person_id=_cell(row, id_col) or f"ps-{index}",
                canonical_name=name,
                aliases=split_eitan_terms(_cell(row, alias_col)),
                role=_cell(row, role_col),
                institution=_cell(row, inst_col),
                view_cluster=_cell(row, cluster_col),
                subcluster=_cell(row, sub_col),
                topic_keywords=split_eitan_terms(_cell(row, topic_col)),
                priority=_priority(_cell(row, pri_col) or 1),
                source_scope=_cell(row, scope_col),
                search_query=_cell(row, query_col),
            )
        )
    return records


def _parse_legacy_sheet(sheet_name: str, rows: list[list[Any]]) -> tuple[list[KeywordRecord], list[PersonRecord], bool]:
    if not rows:
        return [], [], False
    header_kinds = [_classify_header(str(cell or "")) for cell in rows[0]]
    has_header = any(kind for kind in header_kinds)
    body = rows[1:] if has_header else rows
    role = _sheet_kind(sheet_name, _header_map(rows[0]) if rows else {})
    if not has_header:
        if role == "people":
            header_kinds = ["person"]
        elif role == "keywords":
            header_kinds = ["keyword"]
        else:
            header_kinds = (
                ["keyword", "person"]
                if any(len(row) > 1 and _clean_term(row[1]) for row in body[:12])
                else ["keyword"]
            )
    keyword_records: list[KeywordRecord] = []
    person_records: list[PersonRecord] = []
    paired_rows = 0
    filled_rows = 0
    for index, row in enumerate(body, start=1):
        record = {"keyword": [], "person": [], "cluster": "", "weight": 1, "match": ""}
        for col_index, kind in enumerate(header_kinds):
            if not kind or col_index >= len(row):
                continue
            value = row[col_index]
            if kind == "cluster":
                record["cluster"] = _clean_term(value)
            elif kind == "weight":
                record["weight"] = _priority(value)
            elif kind == "match":
                record["match"] = str(value or "")
            elif kind == "keyword":
                record["keyword"].extend(split_eitan_terms(value))
            elif kind == "person":
                record["person"].extend(split_eitan_terms(value))
        keyword_terms = _unique_terms(record["keyword"])
        person_terms = _unique_terms(record["person"])
        if not keyword_terms and not person_terms:
            continue
        filled_rows += 1
        cluster = str(record["cluster"] or "")
        both = bool(keyword_terms and person_terms)
        if both:
            paired_rows += 1
        for offset, term in enumerate(keyword_terms, start=1):
            keyword_records.append(
                KeywordRecord(
                    keyword_id=f"kw-{index}-{offset}",
                    canonical_term=term,
                    category=cluster,
                    priority=int(record["weight"] or 1),
                )
            )
        for offset, term in enumerate(person_terms, start=1):
            person_records.append(
                PersonRecord(
                    person_id=f"ps-{index}-{offset}",
                    canonical_name=term,
                    view_cluster=cluster,
                    priority=int(record["weight"] or 1),
                )
            )
    paired = filled_rows > 0 and (paired_rows / filled_rows) >= 0.3
    return keyword_records, person_records, paired


def parse_eitan_library_upload(filename: str | None, data: bytes) -> EitanLibrary:
    payload = data or b""
    if len(payload) > EITAN_FILE_MAX_BYTES:
        raise ValueError("حجم فایل نباید بیشتر از ۵ مگابایت باشد.")
    name = str(filename or "").strip()
    lowered = name.lower()
    if lowered.endswith((".xlsx", ".xlsm")):
        sheet_names, sheets = _rows_from_xlsx(payload)
    elif lowered.endswith(".csv"):
        sheet_names, sheets = _rows_from_csv(payload)
    else:
        sheet_names, sheets = _rows_from_text(payload)

    keyword_records: list[KeywordRecord] = []
    person_records: list[PersonRecord] = []
    used_structured = False
    paired = False

    for sheet_name, rows in sheets:
        if not rows:
            continue
        headers = _header_map(rows[0])
        kind = _sheet_kind(sheet_name, headers)
        if kind == "keywords":
            parsed = _parse_keyword_sheet(rows)
            if parsed:
                keyword_records.extend(parsed)
                used_structured = True
                continue
        if kind == "people":
            parsed = _parse_person_sheet(rows)
            if parsed:
                person_records.extend(parsed)
                used_structured = True
                continue
        if headers and (headers.keys() & _keys("canonical_term")):
            parsed = _parse_keyword_sheet(rows)
            if parsed:
                keyword_records.extend(parsed)
                used_structured = True
                continue
        if headers and (headers.keys() & _keys("canonical_name")):
            parsed = _parse_person_sheet(rows)
            if parsed:
                person_records.extend(parsed)
                used_structured = True
                continue

    if not used_structured:
        for sheet_name, rows in sheets:
            legacy_keywords, legacy_people, legacy_paired = _parse_legacy_sheet(sheet_name, rows)
            keyword_records.extend(legacy_keywords)
            person_records.extend(legacy_people)
            paired = paired or legacy_paired

    if not keyword_records and not person_records:
        raise ValueError("شیت‌های Keyword_Library یا Person_Library با ستون‌های استاندارد پیدا نشد.")
    return _library_from_records(
        keyword_records,
        person_records,
        filename=name,
        sheets=sheet_names,
        paired=paired,
    )


def parse_eitan_upload(filename: str | None, data: bytes) -> str:
    """Legacy single-column helper used by older two-file uploads."""

    library = parse_eitan_library_upload(filename, data)
    stored = "\n".join(library.keywords or library.people)
    if len(stored) > EITAN_TEXT_MAX_CHARS:
        raise ValueError("حجم فهرست عبارات بیش از حد مجاز است.")
    return stored


def eitan_search_groups(library: EitanLibrary) -> list[list[str]]:
    groups: list[list[str]] = []
    seen: set[tuple[str, ...]] = set()
    used = 0

    def add_group(terms: list[str]) -> None:
        nonlocal used
        cleaned = _unique_terms(terms)[: max(EITAN_AND_WIDTH, 1)]
        if not cleaned:
            return
        key = tuple(canonical_key(term) or term for term in cleaned)
        if key in seen:
            return
        if len(groups) >= EITAN_SEARCH_GROUP_LIMIT or used + len(cleaned) > EITAN_SEARCH_TERM_LIMIT:
            return
        seen.add(key)
        groups.append(cleaned)
        used += len(cleaned)

    source_clauses = library.clauses or _clauses_from_records(library.keyword_records, library.person_records)
    for clause in source_clauses:
        keywords = _unique_terms(clause.get("keywords") or [])
        people = _unique_terms(clause.get("people") or [])
        match = str(clause.get("match") or "any")
        if match == "all" and (keywords or people):
            add_group([*keywords, *people])
            continue
        for term in [*keywords, *people]:
            add_group([term])
        if len(groups) >= EITAN_SEARCH_GROUP_LIMIT:
            break
    if not groups:
        for term in [*library.keywords, *library.people]:
            add_group([term])
            if len(groups) >= EITAN_SEARCH_GROUP_LIMIT:
                break
    return groups[:EITAN_SEARCH_GROUP_LIMIT]


def eitan_chart_terms(library: EitanLibrary, *, limit: int = 24) -> list[dict[str, Any]]:
    ranked: list[tuple[float, str, str]] = []
    for record in library.keyword_records:
        ranked.append((float(record.priority), record.canonical_term, record.category or "keyword"))
    for record in library.person_records:
        ranked.append((float(record.priority), record.canonical_name, record.view_cluster or "person"))
    if not ranked:
        for clause in library.clauses:
            weight = float(clause.get("weight") or 1.0)
            cluster = str(clause.get("cluster") or "")
            for term in _unique_terms([*(clause.get("keywords") or []), *(clause.get("people") or [])]):
                kind = "person" if term in library.people and term not in library.keywords else "keyword"
                ranked.append((weight, term, kind if not cluster else cluster))
    if not ranked:
        ranked = [(1.0, term, "keyword") for term in library.keywords] + [
            (1.0, term, "person") for term in library.people
        ]
    seen: set[str] = set()
    items: list[dict[str, Any]] = []
    for weight, term, kind in sorted(ranked, key=lambda item: (-item[0], item[1])):
        key = canonical_key(term) or term
        if key in seen:
            continue
        seen.add(key)
        items.append({"term": term, "weight": weight, "kind": kind})
        if len(items) >= limit:
            break
    return items


def _term_hits(haystack: str, terms: list[str]) -> bool:
    return any(term and term in haystack for term in terms)


def _ranked(counter: Counter[str], *, limit: int = 12) -> list[dict[str, Any]]:
    return [{"label": label, "count": count} for label, count in counter.most_common(limit)]


def _flow_rows(counter: Counter[tuple[str, str]], *, limit: int = 18) -> list[dict[str, Any]]:
    return [
        {"source": left, "target": right, "count": count}
        for (left, right), count in counter.most_common(limit)
        if left and right
    ]


def build_eitan_insights(library: EitanLibrary, rows: list[dict[str, Any]]) -> dict[str, Any]:
    daily: Counter[str] = Counter()
    keyword_daily: Counter[str] = Counter()
    person_daily: Counter[str] = Counter()
    category_daily: dict[str, Counter[str]] = defaultdict(Counter)
    sources: Counter[str] = Counter()
    message_types: Counter[str] = Counter()
    match_kinds: Counter[str] = Counter()
    keyword_hits: Counter[str] = Counter()
    person_hits: Counter[str] = Counter()
    categories: Counter[str] = Counter()
    subcategories: Counter[str] = Counter()
    heatmap: Counter[tuple[str, str]] = Counter()
    keyword_flow: Counter[tuple[str, str]] = Counter()
    clusters: Counter[str] = Counter()
    subclusters: Counter[str] = Counter()
    institutions: Counter[str] = Counter()
    roles: Counter[str] = Counter()
    cluster_flow: Counter[tuple[str, str]] = Counter()
    person_flow: Counter[tuple[str, str]] = Counter()

    keyword_term_cache = [(record, [normalize_persian(term).casefold() for term in record.search_terms()]) for record in library.keyword_records]
    person_term_cache = [(record, [normalize_persian(term).casefold() for term in record.search_terms()]) for record in library.person_records]

    for row in rows:
        haystack = normalize_persian(str(row.get("haystack") or "")).casefold()
        day = str(row.get("day") or "نامشخص")
        source = str(row.get("source") or "منبع نامشخص")
        message_type = str(row.get("message_type") or "text")
        matched_keyword = False
        matched_person = False
        for record, terms in keyword_term_cache:
            if not _term_hits(haystack, terms):
                continue
            matched_keyword = True
            keyword_hits[record.canonical_term] += 1
            category = record.category or "بدون دسته"
            subcategory = record.subcategory or "بدون زیردسته"
            categories[category] += 1
            subcategories[subcategory] += 1
            heatmap[(category, subcategory)] += 1
            if record.category and record.subcategory:
                keyword_flow[(record.category, record.subcategory)] += 1
            category_daily[category][day] += 1
        for record, terms in person_term_cache:
            if not _term_hits(haystack, terms):
                continue
            matched_person = True
            person_hits[record.canonical_name] += 1
            cluster = record.view_cluster or "بدون خوشه"
            subcluster = record.subcluster or "بدون زیرخوشه"
            clusters[cluster] += 1
            subclusters[subcluster] += 1
            if record.institution:
                institutions[record.institution] += 1
            if record.role:
                roles[record.role] += 1
            if record.view_cluster and record.subcluster:
                cluster_flow[(record.view_cluster, record.subcluster)] += 1
            if record.view_cluster:
                person_flow[(record.view_cluster, record.canonical_name)] += 1
        if matched_keyword or matched_person:
            daily[day] += 1
            sources[source] += 1
            message_types[message_type] += 1
        if matched_keyword:
            keyword_daily[day] += 1
            match_kinds["کلیدواژه"] += 1
        if matched_person:
            person_daily[day] += 1
            match_kinds["افراد شاخص"] += 1

    days = sorted(set(daily) | set(keyword_daily) | set(person_daily))
    top_categories = [label for label, _count in categories.most_common(4)]
    trend_series = [
        {"name": "کل پیام‌ها", "values": [daily.get(day, 0) for day in days], "total": sum(daily.values())},
        {
            "name": "کلیدواژه",
            "values": [keyword_daily.get(day, 0) for day in days],
            "total": sum(keyword_daily.values()),
        },
        {
            "name": "افراد",
            "values": [person_daily.get(day, 0) for day in days],
            "total": sum(person_daily.values()),
        },
    ]
    for category in top_categories:
        series_counts = category_daily.get(category) or Counter()
        trend_series.append(
            {
                "name": category,
                "values": [series_counts.get(day, 0) for day in days],
                "total": sum(series_counts.values()),
            }
        )

    heatmap_rows = [label for label, _count in categories.most_common(10)]
    heatmap_cols = [label for label, _count in subcategories.most_common(10)]
    heatmap_matrix = [
        {
            "row": row_label,
            "cells": [{"col": col_label, "count": int(heatmap.get((row_label, col_label), 0))} for col_label in heatmap_cols],
        }
        for row_label in heatmap_rows
    ]

    keyword_meta = {record.canonical_term: record for record in library.keyword_records}
    person_meta = {record.canonical_name: record for record in library.person_records}
    keyword_rows = []
    for term, count in keyword_hits.most_common(16):
        record = keyword_meta.get(term)
        keyword_rows.append(
            {
                "label": term,
                "count": count,
                "kind": "keyword",
                "category": getattr(record, "category", ""),
                "subcategory": getattr(record, "subcategory", ""),
                "priority": getattr(record, "priority", 1),
            }
        )
    person_rows = []
    for name, count in person_hits.most_common(16):
        record = person_meta.get(name)
        person_rows.append(
            {
                "label": name,
                "count": count,
                "kind": "person",
                "role": getattr(record, "role", ""),
                "institution": getattr(record, "institution", ""),
                "view_cluster": getattr(record, "view_cluster", ""),
                "subcluster": getattr(record, "subcluster", ""),
                "priority": getattr(record, "priority", 1),
            }
        )

    return {
        "sample_size": len(rows),
        "matched_count": int(sum(daily.values())),
        "library": library.summary(),
        "daily": [{"label": day, "count": daily[day]} for day in days],
        "keyword_daily": [{"label": day, "count": keyword_daily.get(day, 0)} for day in days],
        "person_daily": [{"label": day, "count": person_daily.get(day, 0)} for day in days],
        "trend": {"days": days, "series": trend_series},
        "sources": _ranked(sources, limit=10),
        "terms": keyword_rows[:12],
        "keywords": keyword_rows,
        "people": person_rows,
        "kinds": _ranked(match_kinds, limit=8),
        "message_types": _ranked(message_types, limit=8),
        "clusters": _ranked(clusters, limit=12),
        "categories": _ranked(categories, limit=16),
        "subcategories": _ranked(subcategories, limit=12),
        "subclusters": _ranked(subclusters, limit=12),
        "institutions": _ranked(institutions, limit=10),
        "roles": _ranked(roles, limit=10),
        "heatmap": {"rows": heatmap_rows, "cols": heatmap_cols, "matrix": heatmap_matrix},
        "keyword_flow": _flow_rows(keyword_flow),
        "cluster_flow": _flow_rows(cluster_flow),
        "person_flow": _flow_rows(person_flow),
    }
