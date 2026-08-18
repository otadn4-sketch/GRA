from __future__ import annotations

"""Word companion for the HTML-first Garaye layout engine.

The DOCX follows the same publication rules as HTML/PDF: section order,
cover calendars, event cards without photo/title/one-liners, weekday palette,
right-hand news rules, topic-share chart, and editorial footnotes.
"""

from collections import Counter
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Cm, Mm, Pt, RGBColor

from .bulletin_models import BulletinData, BulletinPerson, BulletinStatement, is_event_category
from .bulletin_validation import is_generic_url, valid_public_url
from .calendar_dates import report_calendar_labels
from .html_layout_engine import _category_color, _ordered_public_cards, _topic_share_rows
from .persian_text import to_persian_digits
from .weekday_palette import weekday_palette_for_report

_FOOTNOTES_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:type="continuationSeparator" w:id="0">'
    '<w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
    "</w:footnotes>"
)


def _hex_rgb(value: str) -> RGBColor:
    text = str(value or "#20252B").lstrip("#")
    if len(text) != 6:
        text = "20252B"
    return RGBColor(int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16))


def _rtl(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(4)
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:bidi")) is None:
        p_pr.append(OxmlElement("w:bidi"))


def _set_run(run, *, size: int = 12, bold: bool = False, color: str = "#20252B", font: str = "IRZar") -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = _hex_rgb(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), font)


def _add_text(doc: Document, text: str, *, size: int = 12, bold: bool = False, color: str = "#20252B", center: bool = False, font: str = "IRZar"):
    paragraph = doc.add_paragraph()
    _rtl(paragraph)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(to_persian_digits(text))
    _set_run(run, size=size, bold=bold, color=color, font=font)
    return paragraph


def _shade(cell, color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:shd"))
    if existing is not None:
        tc_pr.remove(existing)
    shd = OxmlElement("w:shd")
    fill = str(color or "#FFFFFF").lstrip("#")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _set_cell_border(cell, *, edge: str = "right", color: str, size: str = "24") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), size)
    element.set(qn("w:space"), "0")
    element.set(qn("w:color"), str(color or "B85C4A").lstrip("#"))
    borders.append(element)


def _set_document_background(doc: Document, color: str) -> None:
    try:
        background = OxmlElement("w:background")
        background.set(qn("w:color"), str(color or "FAF3F1").lstrip("#"))
        root = doc.element
        existing = root.find(qn("w:background"))
        if existing is not None:
            root.remove(existing)
        root.insert(0, background)
        settings = doc.settings.element
        if settings.find(qn("w:displayBackgroundShape")) is None:
            settings.append(OxmlElement("w:displayBackgroundShape"))
    except Exception:
        return


def _set_two_columns(section) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:sep"), "0")
    cols.set(qn("w:space"), "400")


def _set_one_column(section) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), "1")
    cols.set(qn("w:sep"), "0")


def _configure_a3(section, palette: dict[str, Any]) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(297)
    section.page_height = Mm(420)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(22)
    section.left_margin = Mm(16)
    section.right_margin = Mm(16)
    sect_pr = section._sectPr
    pg_sz = sect_pr.find(qn("w:pgSz"))
    if pg_sz is None:
        pg_sz = OxmlElement("w:pgSz")
        sect_pr.append(pg_sz)
    pg_sz.set(qn("w:w"), str(int(round(Mm(297).twips))))
    pg_sz.set(qn("w:h"), str(int(round(Mm(420).twips))))
    pg_sz.set(qn("w:orient"), "portrait")
    pg_mar = sect_pr.find(qn("w:pgMar"))
    if pg_mar is not None:
        pg_mar.set(qn("w:footer"), str(int(round(Mm(14).twips))))
    _ = palette


def _footnotes_part(doc: Document) -> Part:
    for rel in doc.part.rels.values():
        if rel.reltype == RT.FOOTNOTES:
            return rel.target_part
    part = Part(
        PackURI("/word/footnotes.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        _FOOTNOTES_XML.encode("utf-8"),
        doc.part.package,
    )
    doc.part.relate_to(part, RT.FOOTNOTES)
    return part


def _next_footnote_id(footnotes_el) -> int:
    ids = [int(node.get(qn("w:id"))) for node in footnotes_el.findall(qn("w:footnote")) if node.get(qn("w:id"))]
    return max(ids, default=0) + 1


def _add_footnote(doc: Document, paragraph, text: str, color: str) -> int:
    try:
        part = _footnotes_part(doc)
        footnotes_el = getattr(part, "_garaye_footnotes_el", None)
        if footnotes_el is None:
            from docx.oxml import parse_xml

            footnotes_el = parse_xml(part.blob)
            part._garaye_footnotes_el = footnotes_el
        footnote_id = _next_footnote_id(footnotes_el)
        footnote = OxmlElement("w:footnote")
        footnote.set(qn("w:id"), str(footnote_id))
        body = OxmlElement("w:p")
        marker_run = OxmlElement("w:r")
        marker_pr = OxmlElement("w:rPr")
        marker_style = OxmlElement("w:rStyle")
        marker_style.set(qn("w:val"), "FootnoteReference")
        marker_pr.append(marker_style)
        marker_run.append(marker_pr)
        marker_run.append(OxmlElement("w:footnoteRef"))
        body.append(marker_run)
        text_run = OxmlElement("w:r")
        text_pr = OxmlElement("w:rPr")
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "16")
        text_pr.append(size)
        color_el = OxmlElement("w:color")
        color_el.set(qn("w:val"), str(color or "5F6872").lstrip("#"))
        text_pr.append(color_el)
        text_run.append(text_pr)
        text_el = OxmlElement("w:t")
        text_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        text_el.text = f" {to_persian_digits(text)}"
        text_run.append(text_el)
        body.append(text_run)
        footnote.append(body)
        footnotes_el.append(footnote)
        xml = footnotes_el.xml if hasattr(footnotes_el, "xml") else None
        if xml:
            part._blob = xml.encode("utf-8")
        else:
            from lxml import etree

            part._blob = etree.tostring(footnotes_el, xml_declaration=True, encoding="UTF-8", standalone=True)

        run = paragraph.add_run()
        r_pr = run._element.get_or_add_rPr()
        style = OxmlElement("w:rStyle")
        style.set(qn("w:val"), "FootnoteReference")
        vert = OxmlElement("w:vertAlign")
        vert.set(qn("w:val"), "superscript")
        r_pr.append(style)
        r_pr.append(vert)
        ref = OxmlElement("w:footnoteReference")
        ref.set(qn("w:id"), str(footnote_id))
        run._element.append(ref)
        return footnote_id
    except Exception:
        run = paragraph.add_run(to_persian_digits("*"))
        run.font.superscript = True
        run.font.size = Pt(8)
        return 0


def _portrait_bytes(path_value: str | None) -> BytesIO | None:
    if not path_value:
        return None
    path = Path(path_value)
    if not path.is_file():
        return None
    return BytesIO(path.read_bytes())


def _qr_path(statement: BulletinStatement) -> Path | None:
    if not statement.show_qr_in_bulletin:
        return None
    if not valid_public_url(statement.editorial_source_url) or is_generic_url(statement.editorial_source_url):
        return None
    path = Path(statement.qr_code_path or "")
    return path if path.is_file() else None


def _add_news_card(
    doc: Document,
    person: BulletinPerson,
    statement: BulletinStatement,
    *,
    palette: dict[str, Any],
    is_event: bool,
    footnote_text: str = "",
) -> None:
    identity = palette.get("main", "#B85C4A")
    dark = palette.get("dark", "#783B31")
    ink = palette.get("ink", "#20252B")
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    cell = table.cell(0, 0)
    _shade(cell, palette.get("ground", "#FAF3F1"))
    _set_cell_border(cell, edge="right", color=identity, size="24")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    cell.width = Cm(12.4)
    paragraph = cell.paragraphs[0]
    _rtl(paragraph)
    if not is_event:
        portrait = _portrait_bytes(person.portrait_path)
        if portrait:
            try:
                run = paragraph.add_run()
                run.add_picture(portrait, width=Cm(1.6), height=Cm(1.6))
                paragraph.add_run("  ")
            except Exception:
                pass
        name_run = paragraph.add_run(to_persian_digits(person.name_canonical))
        _set_run(name_run, size=12, bold=True, color=dark)
        if person.position:
            paragraph.add_run("\n")
            pos = paragraph.add_run(to_persian_digits(person.position))
            _set_run(pos, size=10, color=palette.get("muted", "#5F6872"))
        headline = statement.main_subject or statement.headline or statement.topic
        if headline:
            h = cell.add_paragraph()
            _rtl(h)
            run = h.add_run(to_persian_digits(headline))
            _set_run(run, size=14, bold=True, color=dark)
        one_line = statement.summary_lead or statement.summary_short
        if one_line:
            line = cell.add_paragraph()
            _rtl(line)
            run = line.add_run(to_persian_digits(one_line))
            _set_run(run, size=11, bold=True, color=ink, font="IRZAR-BOLD")
        location = statement.editorial_context_label or statement.statement_mode
        if location and location not in {"محل بیان نامشخص", "نامشخص"}:
            loc = cell.add_paragraph()
            _rtl(loc)
            run = loc.add_run(to_persian_digits(location))
            _set_run(run, size=10, bold=True, color=dark)
        body_p = cell.add_paragraph()
    else:
        body_p = paragraph
    body = statement.summary_body or statement.summary_detailed or statement.summary_short
    _rtl(body_p)
    run = body_p.add_run(to_persian_digits(body))
    _set_run(run, size=11, color=ink)
    if footnote_text:
        _add_footnote(doc, body_p, footnote_text, palette.get("muted", "#5F6872"))
    qr = _qr_path(statement)
    if qr:
        qr_p = cell.add_paragraph()
        qr_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        try:
            qr_p.add_run().add_picture(str(qr), width=Cm(1.2), height=Cm(1.2))
        except Exception:
            pass
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)


def build_layout_docx(data: BulletinData, path: Path) -> list[str]:
    warnings: list[str] = []
    palette = weekday_palette_for_report(data.meta.report_date_jalali)
    dates = report_calendar_labels(data.meta.report_date_jalali)
    ordered_cards, _ = _ordered_public_cards(data)
    if not ordered_cards:
        raise RuntimeError("برای صفحه‌آرایی بولتن، خبر نهایی قابل انتشار وجود ندارد.")

    doc = Document()
    font = "IRZar"
    identity = palette.get("main", "#B85C4A")
    dark = palette.get("dark", "#783B31")
    ink = palette.get("ink", "#20252B")
    ground = palette.get("ground", "#FAF3F1")
    _set_document_background(doc, ground)
    section = doc.sections[0]
    _configure_a3(section, palette)
    _set_one_column(section)

    cover = doc.add_table(rows=1, cols=1)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_cell = cover.cell(0, 0)
    _shade(cover_cell, dark)
    cover_cell.width = Mm(265)
    p = cover_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("بولتن تحلیلی گرایه")
    _set_run(run, size=14, bold=True, color=palette.get("light", "#EBC4BA"), font=font)
    title = cover_cell.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("خبرنامه گرایه")
    _set_run(run, size=36, bold=True, color="#FFFFFF", font=font)
    copy = cover_cell.add_paragraph()
    copy.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = copy.add_run("رصد، تحلیل و صورت‌بندی هوشمند جریان خبر")
    _set_run(run, size=13, color="#F7F1EF", font=font)
    meta = cover_cell.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(to_persian_digits(f"شماره {data.meta.issue_number}  ·  {dates['jalali_label']}"))
    _set_run(run, size=12, color="#F3E6E2", font=font)
    calendars = cover_cell.add_paragraph()
    calendars.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = calendars.add_run(to_persian_digits(f"قمری {dates['hijri_label']}  ·  میلادی {dates['gregorian_label']}"))
    _set_run(run, size=11, color=palette.get("light", "#EBC4BA"), font=font)
    doc.add_page_break()

    _add_text(doc, "فهرست", size=22, bold=True, color=dark, center=True)
    if data.meta.introduction.strip():
        _add_text(doc, "مقدمه", size=13, color=ink)
    if data.controversies:
        _add_text(doc, "پربازتاب‌ها", size=13, color=ink)
    seen_titles: set[str] = set()
    for category, _person, _statement in ordered_cards:
        if category.title not in seen_titles:
            seen_titles.add(category.title)
            letter = getattr(category, "section_letter", "")
            label = f"{letter}) {category.title}" if letter else category.title
            _add_text(doc, label, size=13, color=ink)
    doc.add_page_break()

    if data.meta.introduction.strip():
        _add_text(doc, "مقدمه", size=22, bold=True, color=dark, center=True)
        for paragraph in str(data.meta.introduction).splitlines():
            if paragraph.strip():
                _add_text(doc, paragraph.strip(), size=13, color=ink)
        doc.add_page_break()

    if data.controversies:
        _add_text(doc, "پربازتاب‌ها", size=22, bold=True, color=dark, center=True)
        for index, item in enumerate(data.controversies, start=1):
            _add_text(doc, f"{index}. {item.title}", size=14, bold=True, color=dark)
            _add_text(doc, item.summary, size=12, color=ink)
        doc.add_page_break()

    content = doc.add_section()
    _configure_a3(content, palette)
    _set_two_columns(content)

    current_category = None
    for category, person, statement in ordered_cards:
        if current_category != category.category_id:
            if current_category is not None:
                doc.add_page_break()
            current_category = category.category_id
            heading = doc.add_paragraph()
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = heading.add_run(to_persian_digits(category.title))
            _set_run(run, size=20, bold=True, color=dark)
        _add_news_card(
            doc,
            person,
            statement,
            palette=palette,
            is_event=is_event_category(category.category_id) or str(category.category_id) == "events",
            footnote_text=str(statement.footnote or "").strip(),
        )

    chart = doc.add_section()
    _configure_a3(chart, palette)
    _set_one_column(chart)
    rows = _topic_share_rows(data)
    if rows:
        _add_text(doc, "سهم موضوعات", size=22, bold=True, color=dark, center=True)
        bar = doc.add_table(rows=1, cols=len(rows))
        bar.alignment = WD_TABLE_ALIGNMENT.CENTER
        for index, item in enumerate(rows):
            cell = bar.cell(0, index)
            _shade(cell, item["color"])
            cell.text = ""
            cell.width = Mm(max(8, 240 * item["share"]))
        for item in rows:
            line = doc.add_paragraph()
            _rtl(line)
            swatch = line.add_run("■ ")
            _set_run(swatch, size=12, color=item["color"])
            label = line.add_run(to_persian_digits(f"{item['topic']}  ·  {item['percent']}٪  ·  {item['count']} خبر"))
            _set_run(label, size=11, color=ink)

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run(to_persian_digits(f"خبرنامه گرایه  ·  شماره {data.meta.issue_number}  ·  {data.meta.report_date_jalali}"))
        _set_run(run, size=9, color=palette.get("muted", "#5F6872"))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    if not path.is_file():
        raise RuntimeError("فایل Word صفحه‌آرایی ساخته نشد.")
    return warnings
