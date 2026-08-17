from __future__ import annotations

import hashlib
import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo

import jdatetime
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .bulletin_builder import BulletinDataBuilder
from .bulletin_models import BulletinData, BulletinStatement, QualityIssue
from .bulletin_validation import has_critical, is_generic_url, valid_public_url, validate_bulletin_data
from .config import Settings
from .db import Database
from .approved_layout_export import ApprovedLayoutExporter
from .html_layout_engine import HtmlLayoutExporter


_STATUS_LABELS = {
    "direct_quote": "نقل مستقیم", "indirect_quote": "نقل غیرمستقیم", "personal_post": "صفحه شخصی",
    "interview": "مصاحبه", "speech": "نشست/سخنرانی", "media_report": "گزارش رسانه‌ای",
}

_ONE_LINE_BOLD_FONT = "IRZAR-BOLD"


def _rtl(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:bidi")) is None:
        p_pr.append(OxmlElement("w:bidi"))


def _set_run_font(run, size: int = 12, bold: bool = False, *, font_name: str = "B Zar") -> None:
    run.font.name = font_name
    run.font.size = Pt(size)
    run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), font_name)


def _set_style_font(style, *, font_name: str, size: int | None = None, bold: bool | None = None) -> None:
    style.font.name = font_name
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), font_name)


def _configure_document_styles(doc: Document, *, font_name: str) -> None:
    styles = {
        "Normal": (12, False), "Title": (24, True), "Subtitle": (15, False),
        "Heading 1": (17, True), "Heading 2": (14, True), "Heading 3": (13, True),
        "Caption": (10, False), "Footer": (9, False), "Header": (9, False),
    }
    for name, (size, bold) in styles.items():
        try:
            _set_style_font(doc.styles[name], font_name=font_name, size=size, bold=bold)
        except KeyError:
            pass
    for style in doc.styles:
        if style.type == WD_STYLE_TYPE.TABLE:
            _set_style_font(style, font_name=font_name, size=10)


def _add_paragraph(doc: Document, text: str = "", *, size: int = 12, bold: bool = False, center: bool = False, font_name: str = "B Zar"):
    p = doc.add_paragraph()
    _rtl(p)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    _set_run_font(run, size=size, bold=bold, font_name=font_name)
    return p


def _add_heading(doc: Document, text: str, level: int, *, font_name: str):
    p = doc.add_heading(text, level=level)
    _rtl(p)
    p.paragraph_format.keep_with_next = True
    for run in p.runs:
        _set_run_font(run, size={1: 17, 2: 14, 3: 13}.get(level, 12), bold=True, font_name=font_name)
    return p


def _add_toc(doc: Document, *, font_name: str) -> None:
    p = doc.add_paragraph()
    _rtl(p)
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "برای نمایش فهرست، در Word گزینه Update Field را انتخاب کنید."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    _set_run_font(run, size=11, font_name=font_name)


def _set_update_fields(doc: Document) -> None:
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def _add_page_number(paragraph, *, font_name: str) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _rtl(paragraph)
    run = paragraph.add_run("صفحه ")
    _set_run_font(run, size=9, font_name=font_name)
    field_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    field_run._r.extend([begin, instr, separate, text, end])
    _set_run_font(field_run, size=9, font_name=font_name)


def _borderless_table(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    bidi = OxmlElement("w:bidiVisual")
    tbl_pr.append(bidi)


def _statement_text(statement: BulletinStatement) -> str:
    mode = (statement.editorial_context_label or statement.statement_mode or "در اظهارنظری").strip().rstrip(":")
    if mode in {"محل بیان نامشخص", "نامشخص", ""}:
        mode = "در اظهارنظری"
    body = statement.summary_body or statement.summary_detailed or statement.summary_lead or statement.summary_short
    return f"{mode}: {body}"


def _add_statement(doc: Document, statement: BulletinStatement, *, font_name: str, qr_size_cm: float, show_qr: bool) -> None:
    text = _statement_text(statement)
    if statement.headline:
        heading = _add_paragraph(doc, statement.headline, size=12, bold=True, font_name=font_name)
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_after = Pt(2)
    one_line = statement.summary_lead or statement.summary_short or statement.headline
    if one_line:
        # This field must remain visibly distinct from the detailed body in
        # every output; Word receives the requested installed font by name.
        lead = _add_paragraph(
            doc,
            "خلاصه: " + one_line,
            size=12,
            bold=True,
            font_name=_ONE_LINE_BOLD_FONT,
        )
        lead.paragraph_format.keep_with_next = True
        lead.paragraph_format.space_after = Pt(2)
    qr_path = Path(statement.qr_code_path) if statement.qr_code_path else None
    if (
        show_qr
        and statement.show_qr_in_bulletin
        and valid_public_url(statement.editorial_source_url)
        and not is_generic_url(statement.editorial_source_url)
        and qr_path
        and qr_path.exists()
    ):
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _borderless_table(table)
        text_cell, qr_cell = table.cell(0, 0), table.cell(0, 1)
        text_cell.width = Cm(14.5)
        qr_cell.width = Cm(max(2.2, qr_size_cm + 0.3))
        text_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        qr_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = text_cell.paragraphs[0]
        _rtl(p)
        p.style = doc.styles["Normal"]
        run = p.add_run("• " + text)
        _set_run_font(run, size=12, font_name=font_name)
        qp = qr_cell.paragraphs[0]
        qp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        qr_run = qp.add_run()
        qr_run.add_picture(str(qr_path), width=Cm(qr_size_cm), height=Cm(qr_size_cm))
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    else:
        p = _add_paragraph(doc, "• " + text, size=12, font_name=font_name)
        p.paragraph_format.space_after = Pt(4)


def _jalali_today(timezone_name: str) -> str:
    local = datetime.now(timezone.utc).astimezone(ZoneInfo(timezone_name))
    jd = jdatetime.datetime.fromgregorian(datetime=local)
    weekdays = ("دوشنبه", "سه‌شنبه", "چهارشنبه", "پنج‌شنبه", "جمعه", "شنبه", "یکشنبه")
    return f"{weekdays[local.weekday()]} {jd.strftime('%Y/%m/%d')}"


def _autosize_sheet(ws, max_width: int = 60) -> None:
    for col in ws.columns:
        width = 10
        for cell in col:
            value = str(cell.value or "")
            width = max(width, min(max_width, len(value) + 2))
            cell.alignment = Alignment(vertical="top", wrap_text=True, readingOrder=2)
        ws.column_dimensions[get_column_letter(col[0].column)].width = width
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", readingOrder=2)


def _write_rows(ws, rows: Iterable[dict[str, Any]], fields: list[str]) -> None:
    ws.append(fields)
    for row in rows:
        values = []
        for field in fields:
            value = row.get(field)
            if isinstance(value, (dict, list, tuple, set)):
                value = json.dumps(value, ensure_ascii=False, default=str)
            values.append(value)
        ws.append(values)
    _autosize_sheet(ws)


class BulletinExporter:
    def __init__(self, db: Database, settings: Settings) -> None:
        self.db = db
        self.settings = settings
        self.builder = BulletinDataBuilder(db, settings)
        project_root = Path(__file__).resolve().parents[1]
        self.html_layout = HtmlLayoutExporter(settings, project_root)
        self.word_layout = ApprovedLayoutExporter(project_root)

    async def export_run(self, run_id: int) -> dict[str, Any]:
        output_dir = self.settings.bulletin_output_dir / f"run_{run_id}"
        output_dir.mkdir(parents=True, exist_ok=True)
        log_id = await self.db.start_bulletin_stage(run_id, "build_json", message="ساخت JSON واسط نسخه ۱۰")
        data, snapshot = await self.builder.build(run_id, output_dir)
        issues = validate_bulletin_data(data)
        if self.settings.bulletin_require_editorial_package and data.meta.editorial_package_status != "validated":
            issues.append(QualityIssue(
                code="editorial_package_missing",
                severity="critical",
                object_type="bulletin_run",
                object_id=str(run_id),
                message=(
                    "بسته بازتدوین سردبیری معتبر موجود نیست یا پس از تصمیم‌های سردبیر قدیمی شده است. "
                    "فقط بازتدوین سردبیری آیتم‌های تأییدشده را اجرا کنید؛ تحلیل کامل پیام‌ها تکرار نمی‌شود."
                ),
                details={"editorial_package_status": data.meta.editorial_package_status},
            ))
        data.quality_control.validation_issues = issues
        await self.db.replace_export_validation_issues(run_id, [x.model_dump() for x in issues])
        await self.db.finish_bulletin_stage(log_id, status="completed_with_warnings" if issues else "completed", level="WARNING" if issues else "INFO", message=f"JSON ساخته شد؛ {len(issues)} مسئله اعتبارسنجی", details={"critical": sum(1 for x in issues if x.severity == 'critical')})

        files: dict[str, str] = {}
        data_path = output_dir / "bulletin_data.json"
        data_path.write_text(data.model_dump_json(indent=2), encoding="utf-8")
        await self._register(run_id, "json", "data", data_path)
        files["data_json"] = str(data_path)

        if self.settings.bulletin_generate_audit_xlsx:
            audit_path = output_dir / "bulletin_audit.xlsx"
            self._build_audit_xlsx(audit_path, data, snapshot, issues)
            await self._register(run_id, "xlsx", "audit", audit_path)
            files["audit_xlsx"] = str(audit_path)

        if self.settings.bulletin_generate_unregistered_xlsx:
            unregistered_path = output_dir / "unregistered_people_review.xlsx"
            self._build_unregistered_xlsx(unregistered_path, data, snapshot)
            await self._register(run_id, "xlsx", "unregistered", unregistered_path)
            files["unregistered_xlsx"] = str(unregistered_path)

        if self.settings.bulletin_generate_run_logs:
            run_log_path = output_dir / "run.log"
            errors_log_path = output_dir / "errors.log"
            self._write_logs(run_log_path, errors_log_path, snapshot, issues)
            await self._register(run_id, "log", "run", run_log_path)
            await self._register(run_id, "log", "errors", errors_log_path)
            files["run_log"] = str(run_log_path)
            files["errors_log"] = str(errors_log_path)

        # Finalization in the editorial desk is the content decision.  This
        # validator stays as an audit report, but it must never stop layout.
        critical = has_critical(issues)
        render_log = await self.db.start_bulletin_stage(
            run_id,
            "render_html_layout",
            message="صفحه‌آرایی مرجع HTML و تولید مستقیم PDF از روی JSON",
        )
        html_export = await self.html_layout.export(data, output_dir)
        layout_html = Path(html_export.files["layout_html"])
        await self._register(run_id, "html", "layout", layout_html)
        files["layout_html"] = str(layout_html)

        layout_report = Path(html_export.files["layout_report"])
        await self._register(run_id, "json", "layout_report", layout_report)
        files["layout_report"] = str(layout_report)

        layout_pdf = html_export.files.get("layout_pdf")
        if layout_pdf:
            pdf_path = Path(layout_pdf)
            await self._register(run_id, "pdf", "classic", pdf_path)
            files["classic_pdf"] = str(pdf_path)

        word_export = self.word_layout.export(data, output_dir)
        layout_docx = Path(word_export.files["layout_docx"])
        await self._register(run_id, "docx", "concise", layout_docx)
        files["concise_docx"] = str(layout_docx)

        layout_warnings = [*html_export.warnings, *word_export.warnings]
        await self.db.finish_bulletin_stage(
            render_log,
            status="completed_with_warnings" if layout_warnings or issues else "completed",
            level="WARNING" if layout_warnings or issues else "INFO",
            message="صفحه‌آرایی مرجع HTML، PDF مستقیم و Word قابل‌ویرایش ساخته شدند",
            details={
                "files": list(files),
                "editorial_version": data.meta.editorial_version,
                "layout_engine": "html-first",
                "layout_audit": html_export.audit,
                "layout_warnings": layout_warnings,
                "validation_is_non_blocking": True,
                "critical_validation_count": sum(1 for issue in issues if issue.severity == "critical"),
            },
        )

        zip_path = output_dir / f"bulletin_run_{run_id}_all.zip"
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for path_str in files.values():
                path = Path(path_str)
                if path.exists():
                    archive.write(path, arcname=path.name)
        await self._register(run_id, "zip", "all", zip_path)
        files["zip"] = str(zip_path)
        return {
            "run_id": run_id,
            "files": files,
            "validation_status": "warning" if issues else "ok",
            "validation_has_critical_issues": critical,
            "word_generated": True,
            "critical_count": sum(1 for x in issues if x.severity == "critical"),
            "warning_count": sum(1 for x in issues if x.severity == "warning"),
        }

    async def _register(self, run_id: int, fmt: str, bucket: str, path: Path) -> None:
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        await self.db.save_bulletin_export(run_id, format_name=fmt, registry_bucket=bucket, file_path=str(path), sha256=digest)

    def _build_docx(self, path: Path, data: BulletinData, *, mode: str) -> None:
        doc = Document()
        font = self.settings.bulletin_font_name or "B Zar"
        _configure_document_styles(doc, font_name=font)
        if self.settings.bulletin_toc_update_on_open:
            _set_update_fields(doc)
        section = doc.sections[0]
        section.top_margin = Cm(1.8); section.bottom_margin = Cm(1.8); section.left_margin = Cm(2.0); section.right_margin = Cm(2.0)

        # Cover
        _add_paragraph(doc, "باسمه تعالی", size=15, bold=True, center=True, font_name=font)
        _add_paragraph(doc, "خبرنامه رصد روزانه", size=25, bold=True, center=True, font_name=font)
        _add_paragraph(doc, "اظهارات شخصیت‌ها", size=20, bold=True, center=True, font_name=font)
        _add_paragraph(doc, f"شماره {data.meta.issue_number}", size=15, bold=True, center=True, font_name=font)
        _add_paragraph(doc, _jalali_today(self.settings.bulletin_default_timezone), size=14, center=True, font_name=font)
        if self.settings.bulletin_show_period_on_cover and data.meta.window_start and data.meta.window_end:
            _add_paragraph(doc, f"بازه رصد: {data.meta.window_start} تا {data.meta.window_end}", size=11, center=True, font_name=font)
        doc.add_page_break()

        _add_heading(doc, "فهرست مطالب", 1, font_name=font)
        _add_toc(doc, font_name=font)
        doc.add_page_break()

        _add_heading(doc, "مقدمه", 1, font_name=font)
        categories = [x.title for x in data.categories]
        intro = (
            f"گزارش پیش رو، شماره {data.meta.issue_number} خبرنامه «اظهارات شخصیت‌ها» است و مهم‌ترین مواضع منتشرشده در بازه "
            f"{data.meta.window_start or 'تعیین‌شده'} تا {data.meta.window_end or 'پایان بازه'} را پوشش می‌دهد. "
            "محتوای رصدشده پس از تشخیص صاحب اظهار، حذف موارد روزمره و تکراری، تفکیک موضوعی و بازبینی سردبیری انتخاب شده است. "
            f"حوزه‌های دارای محتوای قابل انتشار عبارت‌اند از: {('، '.join(categories) if categories else 'حوزه‌های منتخب روز')}. "
            "این خبرنامه گزارش تحلیلی مستقل نیست و مواضع اشخاص را به‌صورت خنثی، فشرده و مستند منعکس می‌کند."
        )
        _add_paragraph(doc, intro, size=12, font_name=font)

        if data.controversies:
            _add_heading(doc, "پربازتاب", 1, font_name=font)
            for index, controversy in enumerate(data.controversies, start=1):
                _add_heading(doc, f"{index}. {controversy.title}", 2, font_name=font)
                _add_paragraph(doc, controversy.summary, font_name=font)
                for side in controversy.sides:
                    _add_paragraph(doc, f"• {side.label}: {side.position_summary}", font_name=font)

        _add_heading(doc, "خلاصه اظهارات", 1, font_name=font)
        for category in data.categories:
            visible_people = [p for p in category.people if any(s.include_in_main for s in p.statements)]
            if not visible_people:
                continue
            _add_heading(doc, f"{category.section_letter}) {category.title}", 1, font_name=font)
            for person in visible_people:
                descriptor = f" - {person.position}" if person.position else ""
                _add_heading(doc, f"{person.continuous_number}. {person.name_canonical}{descriptor}", 2, font_name=font)
                for statement in person.statements:
                    if statement.include_in_main:
                        _add_statement(doc, statement, font_name=font, qr_size_cm=self.settings.bulletin_qr_size_cm, show_qr=self.settings.bulletin_show_qr_codes)

        if mode == "full":
            doc.add_page_break()
            _add_heading(doc, "پیوست الف: مشروح اظهارات", 1, font_name=font)
            for category in data.categories:
                visible_people = [p for p in category.people if any(s.include_in_appendix for s in p.statements)]
                if not visible_people:
                    continue
                _add_heading(doc, category.title, 2, font_name=font)
                for person in visible_people:
                    _add_heading(doc, f"{person.continuous_number}. {person.name_canonical}", 3, font_name=font)
                    for statement in person.statements:
                        if not statement.include_in_appendix:
                            continue
                        _add_paragraph(doc, f"موضوع: {statement.topic}", bold=True, font_name=font)
                        _add_paragraph(doc, f"محل بیان: {statement.editorial_context_label or statement.statement_mode or 'در اظهارنظری'}", size=11, font_name=font)
                        _add_paragraph(doc, statement.summary_detailed, font_name=font)
                        for source in statement.source_records:
                            meta = f"شناسه رکورد: {source.message_id} | منبع: {source.source_name or 'نامشخص'} | زمان: {source.published_at or 'نامشخص'}"
                            if source.source_url:
                                meta += f" | لینک: {source.source_url}"
                            _add_paragraph(doc, meta, size=9, font_name=font)

            _add_heading(doc, "پیوست ب: شخصیت‌های رصدشده", 1, font_name=font)
            for index, person in enumerate(data.monitored_people, start=1):
                status = "داخل شناسنامه" if person.is_registry_person else "خارج از شناسنامه"
                has = "دارای اظهار" if person.has_statement else "فاقد اظهار"
                _add_paragraph(doc, f"{index}. {person.name_canonical} | {person.position or 'سمت نامشخص'} | {person.category or 'دسته نامشخص'} | {status} | {has} | رکورد: {person.record_count} | اظهار نهایی: {person.final_statement_count}", size=10, font_name=font)

            _add_heading(doc, "پیوست ج: منابع رصدی", 1, font_name=font)
            for index, source in enumerate(data.sources, start=1):
                username = f" (@{source.username})" if source.username else ""
                _add_paragraph(doc, f"{index}. {source.title}{username} - {source.record_count} رکورد", size=10, font_name=font)

        for sec in doc.sections:
            footer = sec.footer.paragraphs[0]
            _add_page_number(footer, font_name=font)
        doc.save(path)

    def _build_audit_xlsx(self, path: Path, data: BulletinData, snapshot: dict[str, Any], issues: list[QualityIssue]) -> None:
        wb = Workbook()
        wb.remove(wb.active)
        run = snapshot["run"]
        ws = wb.create_sheet("Run Summary")
        _write_rows(ws, [{
            "run_id": run.get("id"), "issue_number": data.meta.issue_number, "status": run.get("status"),
            "window_start": data.meta.window_start, "window_end": data.meta.window_end,
            "raw_record_count": data.meta.raw_record_count, "accepted_record_count": data.meta.accepted_record_count,
            "approved_item_count": data.meta.approved_item_count, "pipeline_version": data.meta.pipeline_version,
            "validation_critical": sum(1 for x in issues if x.severity == "critical"),
            "validation_warning": sum(1 for x in issues if x.severity == "warning"),
            "provider": run.get("provider"), "model": run.get("model"), "prompt_version": run.get("prompt_version"),
        }], ["run_id","issue_number","status","window_start","window_end","raw_record_count","accepted_record_count","approved_item_count","pipeline_version","validation_critical","validation_warning","provider","model","prompt_version"])

        accepted = []
        clusters = []
        qr_rows = []
        for category in data.categories:
            for person in category.people:
                for st in person.statements:
                    accepted.append({
                        "statement_id": st.statement_id, "person_id": person.person_id, "person_name": person.name_canonical,
                        "category": category.title, "topic": st.topic, "statement_mode": st.statement_mode,
                        "summary_short": st.summary_short, "summary_detailed": st.summary_detailed,
                        "importance_score": st.importance_score, "importance_reason": st.importance_reason,
                        "include_in_main": st.include_in_main, "include_in_appendix": st.include_in_appendix,
                        "confidence": st.confidence, "source_ids": json.dumps(st.source_ids, ensure_ascii=False),
                        "source_url": st.source_url, "qr_code_path": st.qr_code_path,
                    })
                    clusters.append({
                        "statement_id": st.statement_id, "person_name": person.name_canonical, "topic": st.topic,
                        "consensus_method": st.consensus_method, "source_ids": json.dumps(st.source_ids, ensure_ascii=False),
                        "source_count": len(st.source_records), "summary_short": st.summary_short,
                    })
                    qr_rows.append({
                        "statement_id": st.statement_id, "source_url": st.source_url, "has_source_url": st.has_source_url,
                        "qr_code_path": st.qr_code_path, "show_qr": st.show_qr_in_bulletin,
                        "status": "ok" if st.show_qr_in_bulletin else "not_shown",
                    })
        ws = wb.create_sheet("Accepted Statements")
        _write_rows(ws, accepted, list(accepted[0].keys()) if accepted else ["statement_id","person_name","topic"])
        ws = wb.create_sheet("Excluded Records")
        _write_rows(ws, data.quality_control.excluded_records, ["message_id","status","reason"])
        matches = [{
            "message_id": x.get("id"), "detected_person_id": x.get("detected_person_id"), "detected_person_name": x.get("detected_person_name"),
            "method": x.get("person_match_method"), "confidence": x.get("person_confidence"), "matched_alias": x.get("matched_alias"),
            "evidence": x.get("match_context"),
        } for x in snapshot["selected_records"] if x.get("detected_person_name")]
        ws = wb.create_sheet("Person Matches")
        _write_rows(ws, matches, ["message_id","detected_person_id","detected_person_name","method","confidence","matched_alias","evidence"])
        ws = wb.create_sheet("Ambiguous Matches")
        _write_rows(ws, data.quality_control.ambiguous_person_matches, ["message_id","detected_person","confidence","method"])
        ws = wb.create_sheet("Duplicate Groups")
        _write_rows(ws, data.quality_control.duplicate_groups, ["representative_message_id","duplicate_message_ids"])
        ws = wb.create_sheet("Consensus Clusters")
        _write_rows(ws, clusters, list(clusters[0].keys()) if clusters else ["statement_id","person_name","topic"])
        ws = wb.create_sheet("Conflicting Statements")
        _write_rows(ws, data.quality_control.conflicting_statements, ["person_id","topic","statement_ids"])
        ws = wb.create_sheet("Low Confidence")
        _write_rows(ws, data.quality_control.low_confidence_items, ["statement_id","confidence","topic"])
        ws = wb.create_sheet("Controversies")
        controversy_rows = [{
            "title": x.title, "summary": x.summary, "sides": json.dumps([s.model_dump() for s in x.sides], ensure_ascii=False),
            "person_ids": json.dumps(x.person_ids, ensure_ascii=False), "statement_ids": json.dumps(x.statement_ids, ensure_ascii=False),
            "importance_score": x.importance_score, "confidence": x.confidence,
        } for x in data.controversies]
        _write_rows(ws, controversy_rows, ["title","summary","sides","person_ids","statement_ids","importance_score","confidence"])
        ws = wb.create_sheet("Sources")
        _write_rows(ws, [x.model_dump() for x in data.sources], ["source_id","title","username","platform","record_count"])
        ws = wb.create_sheet("QR Validation")
        _write_rows(ws, qr_rows, ["statement_id","source_url","has_source_url","qr_code_path","show_qr","status"])
        ws = wb.create_sheet("Validation Issues")
        validation_rows = [x.model_dump() | {"details": json.dumps(x.details, ensure_ascii=False)} for x in issues]
        _write_rows(ws, validation_rows, ["code","severity","object_type","object_id","message","details"])

        # گزارش امن فراخوانی‌های هوش: فقط اطلاعات عملیاتی و جایگاه کلید ثبت می‌شود؛
        # کلید API، Authorization، متن کامل درخواست و پاسخ عمداً وارد فایل ممیزی نمی‌شوند.
        ai_rows = []
        for request in snapshot.get("ai_requests") or []:
            token_usage = request.get("token_usage")
            if isinstance(token_usage, str):
                try:
                    token_usage = json.loads(token_usage)
                except Exception:
                    token_usage = {"raw": token_usage}
            ai_rows.append({
                "request_id": request.get("request_id"),
                "run_id": request.get("run_id"),
                "item_id": request.get("item_id"),
                "provider": request.get("provider"),
                "model": request.get("model"),
                "prompt_version": request.get("prompt_version"),
                "request_kind": request.get("request_kind"),
                "key_slot": request.get("key_slot"),
                "attempt_number": request.get("attempt_number"),
                "endpoint": request.get("endpoint"),
                "http_status": request.get("http_status"),
                "latency_ms": request.get("latency_ms"),
                "status": request.get("status"),
                "error_text": request.get("error_text"),
                "token_usage": json.dumps(token_usage or {}, ensure_ascii=False),
                "created_at": request.get("created_at"),
            })
        ws = wb.create_sheet("AI Requests")
        _write_rows(
            ws,
            ai_rows,
            [
                "request_id", "run_id", "item_id", "provider", "model", "prompt_version", "request_kind",
                "key_slot", "attempt_number", "endpoint", "http_status", "latency_ms",
                "status", "error_text", "token_usage", "created_at",
            ],
        )
        wb.save(path)

    def _build_unregistered_xlsx(self, path: Path, data: BulletinData, snapshot: dict[str, Any]) -> None:
        wb = Workbook(); ws = wb.active; ws.title = "Unregistered People"
        rows = []
        for category in data.categories:
            for person in category.people:
                if person.is_registry_person:
                    continue
                rows.append({
                    "person_name": person.name_canonical, "category": category.title, "position": person.position,
                    "statement_count": len(person.statements), "topics": " | ".join(x.topic for x in person.statements),
                    "source_message_ids": " | ".join(str(mid) for x in person.statements for mid in x.source_ids),
                    "suggested_action": "بررسی و افزودن/ادغام در شناسنامه",
                })
        _write_rows(ws, rows, ["person_name","category","position","statement_count","topics","source_message_ids","suggested_action"])
        candidates = wb.create_sheet("Person Candidates")
        _write_rows(candidates, snapshot.get("candidates") or [], ["candidate_id","detected_name","confidence","status","sample_message_id","proposed_category","merged_person_id"])
        wb.save(path)

    def _write_logs(self, run_path: Path, error_path: Path, snapshot: dict[str, Any], issues: list[QualityIssue]) -> None:
        run_lines = []
        for row in snapshot.get("logs") or []:
            run_lines.append(f"{row.get('created_at')} | {row.get('level')} | {row.get('stage')} | {row.get('status')} | {row.get('message') or ''} | {row.get('details_json') or ''}")
        run_path.write_text("\n".join(run_lines), encoding="utf-8")
        error_lines = []
        for row in snapshot.get("errors") or []:
            error_lines.append(f"{row.get('created_at')} | {row.get('stage')} | message={row.get('message_id')} | {row.get('error_type')}: {row.get('error_message')}\n{row.get('error_traceback') or ''}")
        for issue in issues:
            if issue.severity in {"warning", "critical"}:
                error_lines.append(f"VALIDATION | {issue.severity} | {issue.code} | {issue.object_type}:{issue.object_id} | {issue.message} | {json.dumps(issue.details, ensure_ascii=False)}")
        error_path.write_text("\n\n".join(error_lines), encoding="utf-8")
